"""Periods, suggestions and decisions through the HTTP API.

The important cases are the refusals: an overlapping window, a second
confirmation of the same movement, approval with pending items, and a stranger
reaching another company's period. A reconciliation feature that only passes the
happy path is a spreadsheet with extra steps.
"""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.deps import get_import_storage
from app.imports.storage import LocalImportStorage
from app.models import AuditLog

pytestmark = pytest.mark.integration

ORGANIZATIONS = "/api/v1/organizations"
BANK_ACCOUNTS = "/api/v1/bank-accounts"
PERIODS = "/api/v1/periods"
MATCHES = "/api/v1/matches"

CSRF_COOKIE = "factuhub_session_csrf"
CSRF_HEADER = "X-CSRF-Token"
PASSWORD = "conciliacion-segura-2026"

MAPPING = {
    "date": "Fecha",
    "debit": "Débito",
    "credit": "Crédito",
    "description": "Descripción",
    "reference": "Referencia",
}

BANK_CSV = (
    "Fecha;Descripción;Débito;Crédito;Referencia\n01/02/2026;Abono cliente;;1.000,00;REF-1\n"
).encode()

LEDGER_CSV = (
    "Fecha;Descripción;Débito;Crédito;Referencia\n01/02/2026;Abono cliente;1.000,00;;REF-1\n"
).encode()


@pytest.fixture(autouse=True)
def isolated_import_storage(client: TestClient, tmp_path) -> LocalImportStorage:
    storage = LocalImportStorage(tmp_path / "imports")
    client.app.dependency_overrides[get_import_storage] = lambda: storage
    return storage


def csrf_headers(client: TestClient) -> dict[str, str]:
    return {CSRF_HEADER: client.cookies[CSRF_COOKIE]}


def company(client: TestClient, tax_id: str = "900123456-7") -> dict[str, Any]:
    response = client.post(
        ORGANIZATIONS,
        json={"name": "Distribuidora El Roble S.A.S.", "tax_id": tax_id},
        headers=csrf_headers(client),
    )
    assert response.status_code == 201, response.text
    return response.json()


def bank_account(client: TestClient, organization_id: str) -> dict[str, Any]:
    response = client.post(
        f"{ORGANIZATIONS}/{organization_id}/bank-accounts",
        json={
            "account_name": "Cuenta corriente",
            "bank_name": "Bancolombia",
            "account_number_masked": "****1234",
            "currency": "COP",
        },
        headers=csrf_headers(client),
    )
    assert response.status_code == 201, response.text
    return response.json()


def upload(
    client: TestClient,
    organization_id: str,
    bank_account_id: str,
    content: bytes,
    *,
    source_type: str,
    filename: str,
) -> dict[str, Any]:
    response = client.post(
        f"{ORGANIZATIONS}/{organization_id}/imports",
        data={
            "source_type": source_type,
            "bank_account_id": bank_account_id,
            "mapping": json.dumps(MAPPING),
        },
        files={"file": (filename, content, "text/csv")},
        headers=csrf_headers(client),
    )
    assert response.status_code == 201, response.text
    return response.json()


def import_pair(
    client: TestClient,
    organization_id: str,
    account_id: str,
    *,
    bank: bytes = BANK_CSV,
    ledger: bytes = LEDGER_CSV,
) -> tuple[dict[str, Any], dict[str, Any]]:
    return (
        upload(
            client,
            organization_id,
            account_id,
            bank,
            source_type="bank_statement",
            filename="extracto.csv",
        ),
        upload(
            client,
            organization_id,
            account_id,
            ledger,
            source_type="ledger_export",
            filename="contable.csv",
        ),
    )


def create_period(
    client: TestClient,
    account_id: str,
    *,
    start: str = "2026-02-01",
    end: str = "2026-02-28",
    opening: int = 0,
    bank_closing: int = 100_000,
    ledger_closing: int = 100_000,
):
    return client.post(
        f"{BANK_ACCOUNTS}/{account_id}/periods",
        json={
            "period_start": start,
            "period_end": end,
            "opening_balance_minor": opening,
            "bank_closing_balance_minor": bank_closing,
            "ledger_closing_balance_minor": ledger_closing,
        },
        headers=csrf_headers(client),
    )


def register(client: TestClient, email: str) -> None:
    client.post(
        "/api/v1/auth/register",
        json={"email": email, "full_name": "Persona de prueba", "password": PASSWORD},
    )


def sign_in(client: TestClient, email: str) -> TestClient:
    register(client, email)
    other = TestClient(client.app)
    other.post("/api/v1/auth/login", json={"email": email, "password": PASSWORD})
    return other


def add_member(client: TestClient, organization_id: str, email: str, role: str) -> None:
    register(client, email)
    response = client.post(
        f"{ORGANIZATIONS}/{organization_id}/members",
        json={"email": email, "role": role},
        headers=csrf_headers(client),
    )
    assert response.status_code == 201, response.text


def generate(client: TestClient, period_id: str) -> dict[str, Any]:
    response = client.post(f"{PERIODS}/{period_id}/match-suggestions", headers=csrf_headers(client))
    assert response.status_code == 200, response.text
    return response.json()


def test_a_period_is_created_and_read(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])

    created = create_period(auth_client, account["id"])

    assert created.status_code == 201, created.text
    body = created.json()
    assert body["status"] == "open"
    assert body["period_start"] == "2026-02-01"
    assert auth_client.get(f"{PERIODS}/{body['id']}").status_code == 200


def test_an_overlapping_window_is_refused(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    assert create_period(auth_client, account["id"]).status_code == 201

    overlap = create_period(auth_client, account["id"])
    adjacent = create_period(auth_client, account["id"], start="2026-03-01", end="2026-03-31")

    assert overlap.status_code == 409
    assert adjacent.status_code == 201


def test_a_reversed_window_is_refused(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])

    response = create_period(auth_client, account["id"], start="2026-02-28", end="2026-02-01")

    assert response.status_code == 422


def test_period_transactions_show_both_sides(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()

    response = auth_client.get(f"{PERIODS}/{period['id']}/transactions")

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["total"] == 2
    assert {row["side"] for row in body["transactions"]} == {"bank", "ledger"}


def test_the_summary_starts_with_pending_movements(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()

    body = auth_client.get(f"{PERIODS}/{period['id']}/summary").json()

    assert body["balances"]["pending_bank_count"] == 1
    assert body["balances"]["pending_ledger_count"] == 1
    assert body["balances"]["progress_ratio"] == 0.0
    assert body["balances"]["balanced"] is False


def test_an_exact_match_is_suggested_confirmed_and_balanced(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()

    suggestions = generate(auth_client, period["id"])
    assert suggestions["created"] == 1
    match = suggestions["matches"][0]
    assert match["match_type"] == "exact"
    assert match["score"] == 1.0
    assert len(match["items"]) == 2

    confirmed = auth_client.post(
        f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client)
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "confirmed"
    assert confirmed.json()["reviewed_by"] is not None

    summary = auth_client.get(f"{PERIODS}/{period['id']}/summary").json()
    assert summary["balances"]["pending_bank_count"] == 0
    assert summary["balances"]["pending_ledger_count"] == 0
    assert summary["balances"]["matched_bank_count"] == 1
    assert summary["balances"]["progress_ratio"] == 1.0
    assert summary["balances"]["balanced"] is True


def test_regenerating_suggestions_replaces_the_previous_set(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()

    first = generate(auth_client, period["id"])
    second = generate(auth_client, period["id"])
    listing = auth_client.get(f"{PERIODS}/{period['id']}/matches").json()

    assert first["created"] == 1
    assert second["created"] == 1
    assert second["matches"][0]["id"] != first["matches"][0]["id"]
    assert len(listing["matches"]) == 1


def test_a_rejected_combination_is_not_offered_again(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]

    rejected = auth_client.post(
        f"{MATCHES}/{match['id']}/reject", headers=csrf_headers(auth_client)
    )
    assert rejected.status_code == 200
    assert rejected.json()["status"] == "rejected"

    again = generate(auth_client, period["id"])

    assert again["created"] == 0


def test_confirming_one_match_cancels_the_alternatives_sharing_a_movement(
    auth_client: TestClient,
) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    bank = (
        "Fecha;Descripción;Débito;Crédito;Referencia\n"
        "01/02/2026;Abono cliente;;1.000,00;REF-1\n"
        "01/02/2026;Abono cliente;;1.000,00;REF-1\n"
    ).encode()
    import_pair(auth_client, organization["id"], account["id"], bank=bank)
    period = create_period(auth_client, account["id"]).json()
    suggestions = generate(auth_client, period["id"])
    assert suggestions["created"] == 2

    first, second = suggestions["matches"]
    confirmed = auth_client.post(
        f"{MATCHES}/{first['id']}/confirm", headers=csrf_headers(auth_client)
    )
    assert confirmed.status_code == 200

    listing = auth_client.get(f"{PERIODS}/{period['id']}/matches").json()
    statuses = {match["id"]: match["status"] for match in listing["matches"]}
    assert statuses[first["id"]] == "confirmed"
    assert statuses[second["id"]] == "cancelled"

    retry = auth_client.post(f"{MATCHES}/{second['id']}/confirm", headers=csrf_headers(auth_client))
    assert retry.status_code == 409


def test_a_completed_match_cannot_be_confirmed_twice(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))

    second = auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))

    assert second.status_code == 409


def test_a_grouped_suggestion_sums_exactly(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    bank = (
        "Fecha;Descripción;Débito;Crédito;Referencia\n01/02/2026;Cobro factura;;1.000,00;FAC000\n"
    ).encode()
    ledger = (
        "Fecha;Descripción;Débito;Crédito;Referencia\n"
        "01/02/2026;Factura A;600,00;;FAC001\n"
        "01/02/2026;Factura B;400,00;;FAC002\n"
    ).encode()
    import_pair(auth_client, organization["id"], account["id"], bank=bank, ledger=ledger)
    period = create_period(
        auth_client, account["id"], bank_closing=100_000, ledger_closing=100_000
    ).json()

    suggestions = generate(auth_client, period["id"])

    assert suggestions["created"] == 1
    match = suggestions["matches"][0]
    assert match["match_type"] == "grouped"
    assert len(match["items"]) == 3

    confirmed = auth_client.post(
        f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client)
    )
    assert confirmed.status_code == 200
    summary = auth_client.get(f"{PERIODS}/{period['id']}/summary").json()
    assert summary["balances"]["balanced"] is True


def test_a_missing_bank_row_blocks_approval(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    # The statement says 1,500 but the imported row only adds 1,000.
    period = create_period(
        auth_client, account["id"], bank_closing=150_000, ledger_closing=150_000
    ).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(auth_client))

    approved = auth_client.post(
        f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client)
    )

    assert approved.status_code == 409
    summary = auth_client.get(f"{PERIODS}/{period['id']}/summary").json()
    assert summary["balances"]["statement_difference_minor"] == 50_000
    assert summary["balances"]["balanced"] is False


def test_a_period_with_pending_movements_cannot_be_approved(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    auth_client.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(auth_client))

    approved = auth_client.post(
        f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client)
    )

    assert approved.status_code == 409
    assert "pending" in approved.json()["detail"]


def test_approval_requires_the_review_state(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    period = create_period(auth_client, account["id"]).json()

    approved = auth_client.post(
        f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client)
    )

    assert approved.status_code == 409


def test_an_approved_period_can_be_reopened(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(auth_client))

    approved = auth_client.post(
        f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client)
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "approved"
    assert approved.json()["approved_by"] is not None

    reopened = auth_client.post(
        f"{PERIODS}/{period['id']}/reopen", headers=csrf_headers(auth_client)
    )
    assert reopened.status_code == 200
    assert reopened.json()["status"] == "reopened"
    assert reopened.json()["approved_by"] is None


def test_a_viewer_can_read_but_not_reconcile(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    add_member(auth_client, organization["id"], "viewer@example.com", "viewer")
    viewer = sign_in(auth_client, "viewer@example.com")

    assert viewer.get(f"{PERIODS}/{period['id']}").status_code == 200
    denied = create_period(viewer, account["id"], start="2026-03-01", end="2026-03-31")
    suggestions = viewer.post(
        f"{PERIODS}/{period['id']}/match-suggestions", headers=csrf_headers(viewer)
    )

    assert denied.status_code == 403
    assert suggestions.status_code == 403


def test_a_reviewer_can_approve_but_not_reconcile(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    add_member(auth_client, organization["id"], "revisor@example.com", "reviewer")
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(auth_client))

    reviewer = sign_in(auth_client, "revisor@example.com")
    denied = create_period(reviewer, account["id"], start="2026-04-01", end="2026-04-30")
    approved = reviewer.post(f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(reviewer))

    assert denied.status_code == 403
    assert approved.status_code == 200, approved.text


def test_a_period_from_another_company_is_a_404(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]

    stranger = sign_in(auth_client, "intruso@example.com")

    assert stranger.get(f"{PERIODS}/{period['id']}").status_code == 404
    assert stranger.get(f"{MATCHES}/{match['id']}").status_code == 404
    assert (
        stranger.post(
            f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(stranger)
        ).status_code
        == 404
    )


def test_an_account_from_another_company_cannot_open_a_period(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    stranger = sign_in(auth_client, "intruso@example.com")

    response = create_period(stranger, account["id"])

    assert response.status_code == 404


def test_an_approved_period_cannot_be_approved_twice(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(auth_client))
    assert (
        auth_client.post(
            f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client)
        ).status_code
        == 200
    )

    second = auth_client.post(
        f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client)
    )

    assert second.status_code == 409


def test_reopening_requires_the_approval_permission(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client))
    add_member(auth_client, organization["id"], "consulta@example.com", "viewer")
    viewer = sign_in(auth_client, "consulta@example.com")

    reopened = viewer.post(f"{PERIODS}/{period['id']}/reopen", headers=csrf_headers(viewer))
    submitted = viewer.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(viewer))

    assert reopened.status_code == 403
    assert submitted.status_code == 403


def test_a_confirmed_match_can_be_cancelled_and_offered_again(
    auth_client: TestClient,
) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))

    cancelled = auth_client.post(
        f"{MATCHES}/{match['id']}/cancel", headers=csrf_headers(auth_client)
    )

    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["status"] == "cancelled"
    assert cancelled.json()["released"] is True
    summary = auth_client.get(f"{PERIODS}/{period['id']}/summary").json()
    assert summary["balances"]["pending_bank_count"] == 1
    assert summary["balances"]["pending_ledger_count"] == 1

    again = generate(auth_client, period["id"])

    assert again["created"] == 1
    assert again["matches"][0]["status"] == "suggested"


def test_undoing_a_confirmation_reoffers_the_cancelled_alternatives(
    auth_client: TestClient,
) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    bank = (
        "Fecha;Descripción;Débito;Crédito;Referencia\n"
        "01/02/2026;Abono cliente;;1.000,00;REF-1\n"
        "01/02/2026;Abono cliente;;1.000,00;REF-1\n"
    ).encode()
    import_pair(auth_client, organization["id"], account["id"], bank=bank)
    period = create_period(auth_client, account["id"]).json()
    first, second = generate(auth_client, period["id"])["matches"]
    auth_client.post(f"{MATCHES}/{first['id']}/confirm", headers=csrf_headers(auth_client))
    after_confirm = auth_client.get(f"{PERIODS}/{period['id']}/matches").json()["matches"]
    assert {row["id"]: row["status"] for row in after_confirm}[second["id"]] == "cancelled"

    auth_client.post(f"{MATCHES}/{first['id']}/cancel", headers=csrf_headers(auth_client))

    again = generate(auth_client, period["id"])

    assert again["created"] == 2


def test_confirming_a_movement_confirmed_elsewhere_is_refused(
    auth_client: TestClient, db_session
) -> None:
    from uuid import UUID

    from app.models import MatchStatus, ReconciliationMatch

    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    bank = (
        "Fecha;Descripción;Débito;Crédito;Referencia\n"
        "01/02/2026;Abono cliente;;1.000,00;REF-1\n"
        "01/02/2026;Abono cliente;;1.000,00;REF-1\n"
    ).encode()
    import_pair(auth_client, organization["id"], account["id"], bank=bank)
    period = create_period(auth_client, account["id"]).json()
    first, second = generate(auth_client, period["id"])["matches"]

    # Simulate a confirmation made elsewhere that the auto-cancellation did not
    # see (for example, a concurrent transaction that committed first).
    row = db_session.get(ReconciliationMatch, UUID(first["id"]))
    assert row is not None
    row.status = MatchStatus.CONFIRMED
    db_session.flush()

    response = auth_client.post(
        f"{MATCHES}/{second['id']}/confirm", headers=csrf_headers(auth_client)
    )

    assert response.status_code == 409
    assert "already confirmed" in response.json()["detail"]


def _summary_value(workbook, label: str):
    sheet = workbook["Resumen"]
    for row in sheet.iter_rows(min_col=1, max_col=2):
        if row[0].value == label:
            return row[1].value
    raise AssertionError(f"label not found in summary: {label}")


def test_the_period_exports_as_an_excel_workbook(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))

    response = auth_client.get(f"{PERIODS}/{period['id']}/export")

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        'filename="conciliacion-2026-02-01-2026-02-28.xlsx"'
        in response.headers["content-disposition"]
    )

    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["Resumen", "Movimientos", "Pendientes", "Coincidencias"]
    assert _summary_value(workbook, "Empresa") == "Distribuidora El Roble S.A.S."
    assert _summary_value(workbook, "Saldo final calculado (banco)") == 1000.0
    assert _summary_value(workbook, "Diferencia no explicada") == 0.0
    assert _summary_value(workbook, "Conciliado") == "Sí"

    movements = workbook["Movimientos"]
    assert movements.max_row == 5  # title, header, bank and ledger rows
    assert movements.cell(row=4, column=1).value == "Banco"
    assert movements.cell(row=4, column=3).value == 1000.0

    matches = workbook["Coincidencias"]
    assert matches.max_row == 4
    assert matches.cell(row=4, column=2).value == "Exacta"
    assert matches.cell(row=4, column=4).value == "Confirmada"

    pending = workbook["Pendientes"]
    assert pending.max_row == 3  # title and header, no pending movement


def test_the_export_keeps_formula_like_text_as_text(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    bank = (
        "Fecha;Descripción;Débito;Crédito;Referencia\n01/02/2026;=1+1;;1.000,00;REF-1\n"
    ).encode()
    import_pair(auth_client, organization["id"], account["id"], bank=bank)
    period = create_period(auth_client, account["id"]).json()

    response = auth_client.get(f"{PERIODS}/{period['id']}/export")

    workbook = load_workbook(BytesIO(response.content), data_only=False)
    sheet = workbook["Movimientos"]
    description = sheet.cell(row=4, column=6)
    assert description.value == "=1+1"
    assert description.data_type == "s"


def test_the_export_records_the_download_in_the_audit_trail(
    auth_client: TestClient, db_session
) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    period = create_period(auth_client, account["id"]).json()

    auth_client.get(f"{PERIODS}/{period['id']}/export")

    entry = db_session.scalar(
        select(AuditLog).where(
            AuditLog.action == "reconciliation_period.exported",
            AuditLog.entity_id == UUID(period["id"]),
        )
    )
    assert entry is not None


def test_a_stranger_cannot_export_a_period(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    period = create_period(auth_client, account["id"]).json()
    stranger = sign_in(auth_client, "intruso@example.com")

    response = stranger.get(f"{PERIODS}/{period['id']}/export")

    assert response.status_code == 404


def test_a_viewer_can_export_the_period(auth_client: TestClient) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    period = create_period(auth_client, account["id"]).json()
    add_member(auth_client, organization["id"], "exporta@example.com", "viewer")
    viewer = sign_in(auth_client, "exporta@example.com")

    response = viewer.get(f"{PERIODS}/{period['id']}/export")

    assert response.status_code == 200


def test_undoing_a_confirmation_requires_reopening_an_approved_period(
    auth_client: TestClient,
) -> None:
    organization = company(auth_client)
    account = bank_account(auth_client, organization["id"])
    import_pair(auth_client, organization["id"], account["id"])
    period = create_period(auth_client, account["id"]).json()
    match = generate(auth_client, period["id"])["matches"][0]
    auth_client.post(f"{MATCHES}/{match['id']}/confirm", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/submit-review", headers=csrf_headers(auth_client))
    auth_client.post(f"{PERIODS}/{period['id']}/approve", headers=csrf_headers(auth_client))

    blocked = auth_client.post(f"{MATCHES}/{match['id']}/cancel", headers=csrf_headers(auth_client))
    assert blocked.status_code == 409
    assert "Reopen" in blocked.json()["detail"]

    auth_client.post(f"{PERIODS}/{period['id']}/reopen", headers=csrf_headers(auth_client))
    released = auth_client.post(
        f"{MATCHES}/{match['id']}/cancel", headers=csrf_headers(auth_client)
    )
    assert released.status_code == 200
    assert released.json()["released"] is True
