"""Reading a CSV or XLSX file into headers and rows.

The reader never trusts the extension alone: an `.xlsx` must really start with a
ZIP signature, a CSV must decode as text. Reading is bounded by `max_rows`, and
`data_only=True` plus a refusal of `.xlsm` keeps a workbook from being anything
more than data.

Cell values stay in their native type where the format has one. An XLSX amount
cell is a number and a date cell is a date; converting them to text first would
reintroduce the separator ambiguity that `parsing.py` exists to avoid.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

type CellValue = str | int | float | date | datetime | time | None

DELIMITERS = ",;\t|"


class ImportFileError(ValueError):
    """The file itself is unusable, before any row is interpreted."""


@dataclass(frozen=True, slots=True)
class TableRow:
    """One data row and the spreadsheet line it came from."""

    number: int
    cells: dict[str, CellValue]


@dataclass(frozen=True, slots=True)
class Table:
    headers: tuple[str, ...]
    rows: tuple[TableRow, ...]
    file_format: str


def read_table(
    data: bytes,
    *,
    file_name: str,
    content_type: str,
    max_rows: int,
) -> Table:
    """Read `data` as a table, refusing anything outside the accepted formats."""
    if not data:
        raise ImportFileError("The file is empty.")

    if "macroenabled" in (content_type or "").lower():
        raise ImportFileError("Macro-enabled workbooks are not accepted; save it as .xlsx.")

    suffix = Path(file_name).suffix.lower()
    if suffix in {".xlsm", ".xltm"}:
        raise ImportFileError("Macro-enabled workbooks are not accepted; save it as .xlsx.")
    if suffix == ".xlsx":
        return _read_xlsx(data, max_rows=max_rows)
    if suffix in {".csv", ".txt"}:
        return _read_csv(data, max_rows=max_rows)
    raise ImportFileError("Only .csv, .txt and .xlsx files are accepted.")


def _read_csv(data: bytes, *, max_rows: int) -> Table:
    text = _decode(data)
    if not text.strip():
        raise ImportFileError("The file has no header row.")
    delimiter = _detect_delimiter(text)

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    headers: list[str] | None = None
    rows: list[TableRow] = []

    for number, values in enumerate(reader, start=1):
        if headers is None:
            if all(_is_blank(value) for value in values):
                continue
            headers = _clean_headers(values, number)
            continue

        if len(values) > len(headers):
            raise ImportFileError(
                f"Row {number} has {len(values)} columns but the header has {len(headers)}."
            )
        cells = {
            header: (values[position].strip() if position < len(values) else "")
            for position, header in enumerate(headers)
        }
        if all(value == "" for value in cells.values()):
            continue
        rows.append(TableRow(number=number, cells=cells))
        if len(rows) > max_rows:
            raise ImportFileError(f"The file has more than {max_rows} data rows.")

    if headers is None:
        raise ImportFileError("The file has no header row.")

    return Table(headers=tuple(headers), rows=tuple(rows), file_format="csv")


def _read_xlsx(data: bytes, *, max_rows: int) -> Table:
    if not data.startswith(b"PK\x03\x04"):
        raise ImportFileError("The .xlsx file is not a valid ZIP container.")

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ImportFileError("The .xlsx file could not be read.") from exc

    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ImportFileError("The workbook has no worksheet.")

        headers: list[str] | None = None
        rows: list[TableRow] = []

        for number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            trimmed = _trim_trailing_empty(values)
            if headers is None:
                if all(_is_blank(value) for value in trimmed):
                    continue
                headers = _clean_headers(trimmed, number)
                continue

            if len(trimmed) > len(headers):
                raise ImportFileError(
                    f"Row {number} has {len(trimmed)} columns but the header has {len(headers)}."
                )
            cells = {
                header: _clean_cell(trimmed[position] if position < len(trimmed) else None)
                for position, header in enumerate(headers)
            }
            if all(_is_blank(value) for value in cells.values()):
                continue
            rows.append(TableRow(number=number, cells=cells))
            if len(rows) > max_rows:
                raise ImportFileError(f"The file has more than {max_rows} data rows.")

        if headers is None:
            raise ImportFileError("The file has no header row.")

        return Table(headers=tuple(headers), rows=tuple(rows), file_format="xlsx")
    finally:
        workbook.close()


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                # Excel's "CSV (semicolon-delimited)" on Spanish Windows is
                # frequently Windows-1252, not UTF-8.
                text = data.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise ImportFileError("The file is not UTF-8 or Windows-1252 text.") from exc

    if "\x00" in text:
        raise ImportFileError("The file is not a text CSV.")
    return text


def _detect_delimiter(text: str) -> str:
    sample = text[:8192]
    lines = sample.splitlines()
    first_line = lines[0] if lines else ""

    # The header decides first. A semicolon file whose amounts use `1.234,56`
    # contains far more commas than semicolons in the body, and a body-wide sniff
    # would pick `,` and split every amount in two.
    counts = {delimiter: first_line.count(delimiter) for delimiter in DELIMITERS}
    chosen = max(counts, key=lambda item: counts[item])
    if counts[chosen] > 0:
        return chosen

    try:
        return csv.Sniffer().sniff(sample, delimiters=DELIMITERS).delimiter
    except csv.Error:
        raise ImportFileError(
            "The CSV has a single column; a delimiter such as ',' or ';' is required."
        ) from None


def _clean_headers(values: list[CellValue] | tuple[CellValue, ...], number: int) -> list[str]:
    """Return the header names, rejecting empty and repeated names."""
    cleaned = [("" if value is None else str(value).strip()) for value in values]
    while cleaned and cleaned[-1] == "":
        cleaned.pop()

    if not cleaned or any(name == "" for name in cleaned):
        raise ImportFileError(f"Row {number} has an empty column name.")
    duplicates = sorted({name for name in cleaned if cleaned.count(name) > 1})
    if duplicates:
        raise ImportFileError(f"The header repeats the column(s): {', '.join(duplicates)}.")
    return cleaned


def _trim_trailing_empty(values: tuple[CellValue, ...]) -> list[CellValue]:
    trimmed = list(values)
    while trimmed and _is_blank(trimmed[-1]):
        trimmed.pop()
    return trimmed


def _clean_cell(value: CellValue) -> CellValue:
    if isinstance(value, str):
        return value.strip()
    return value


def _is_blank(value: CellValue) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
